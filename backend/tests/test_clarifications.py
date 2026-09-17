import io
import json

import pytest
from openpyxl import Workbook, load_workbook

from app.main import app
from app.providers import get_llm_provider
from app.providers.base import LLMUsage


class FakeProvider:
    def __init__(self):
        self.calls = 0

    @property
    def usage(self) -> LLMUsage:
        return LLMUsage(input_tokens=100 * self.calls, output_tokens=20 * self.calls, remaining_tokens=9880)

    async def generate_text(self, prompt: str, system_prompt: str | None = None) -> str:
        self.calls += 1
        if system_prompt and "Duplicate Agent" in system_prompt:
            return json.dumps({"candidates": []})
        if system_prompt and "Sprint Agent" in system_prompt:
            return json.dumps({"decisions": [{
                "story_stable_id": "STORY-001", "decision": "planned",
                "sprint_id": "SPR-1", "assignee_id": "MEM-1",
                "reason": "The high-priority story fits within Sprint 1 capacity and has an available owner.",
                "confidence": 0.88,
            }]})
        if system_prompt and "Assignment Agent" in system_prompt:
            return json.dumps({"assignments": [
                {
                    "item_stable_id": "STORY-001", "team_member_id": "MEM-1",
                    "recommended_hours": None, "match_score": 0.9,
                    "reason": "The engineer has the identity integration skills needed to own this story.",
                    "confidence": 0.87,
                },
                {
                    "item_stable_id": "TASK-001", "team_member_id": "MEM-1",
                    "recommended_hours": 12, "match_score": 0.92,
                    "reason": "The engineer has Python and OIDC integration experience with sufficient capacity.",
                    "confidence": 0.89,
                },
            ]})
        if system_prompt and "Dependency Agent" in system_prompt:
            return json.dumps({"dependencies": []})
        if system_prompt and "Estimation Agent" in system_prompt:
            return json.dumps({
                "stories": [{
                    "stable_id": "STORY-001", "story_points": 5,
                    "rationale": "OIDC integration has moderate complexity and external coordination.",
                    "confidence": 0.82,
                }],
                "tasks": [{
                    "stable_id": "TASK-001", "estimated_hours": 12,
                    "rationale": "Configuration, implementation, and automated tests require twelve hours.",
                    "confidence": 0.78,
                }],
            })
        if system_prompt and "Enrichment Agent" in system_prompt:
            return json.dumps({
                "epics": [{
                    "stable_id": "EPIC-001",
                    "business_value": "Protects project data and reduces unauthorized access risk.",
                    "priority": "Critical",
                    "acceptance_criteria": ["Only authenticated users can access protected project functions."],
                    "confidence": 0.93,
                }],
                "stories": [{
                    "stable_id": "STORY-001", "priority": "High",
                    "acceptance_criteria": ["Given a valid OIDC identity, when login completes, then the dashboard opens."],
                    "confidence": 0.91,
                }],
                "tasks": [{"stable_id": "TASK-001", "priority": "High", "confidence": 0.88}],
            })
        if system_prompt and "Backlog Agent" in system_prompt:
            return json.dumps({
                "epics": [{
                    "epic_key": "E1", "decomposition_ids": ["DEC-001"],
                    "title": "Identity and access management",
                    "description": "Provide secure access to protected project capabilities.",
                    "business_value": "Protects project data while enabling authorized work.",
                    "confidence": 0.9,
                }],
                "stories": [{
                    "story_key": "S1", "parent_epic_key": "E1", "decomposition_ids": ["DEC-001"],
                    "title": "Authenticate with OIDC",
                    "user_story": "As a user, I want to authenticate with OIDC so that I can access the platform securely.",
                    "description": "Authenticate platform users through the approved OIDC identity provider.",
                    "confidence": 0.92,
                }],
                "tasks": [{
                    "task_key": "T1", "parent_story_key": "S1", "decomposition_ids": ["DEC-001"],
                    "title": "Integrate OIDC provider",
                    "description": "Configure and implement the approved OIDC authentication integration.",
                    "task_type": "implementation", "confidence": 0.88,
                }],
            })
        if system_prompt and "Decomposition Agent" in system_prompt:
            return json.dumps({"decompositions": [{
                "requirement_ids": ["REQ-001"],
                "parent_capability": "Identity and access management",
                "component_type": "business_capability",
                "title": "User authentication",
                "description": "Provide secure authentication through the approved identity provider.",
                "suggested_backlog_level": "epic",
                "rationale": "Authentication is a distinct business capability with multiple delivery concerns.",
                "confidence": 0.9,
            }]})
        if system_prompt and "Requirement agent" in system_prompt:
            return json.dumps({"requirements": [{
                "title": "Authenticate users with OIDC",
                "description": "The platform must authenticate every user through the approved OIDC identity provider.",
                "category": "functional",
                "requirement_status": "explicit",
                "actors": ["User"],
                "systems": ["OIDC identity provider"],
                "business_rules": [],
                "constraints": [],
                "assumptions": [],
                "priority": "High",
                "rationale": "Authentication protects access to project information.",
                "acceptance_criteria": ["Unauthenticated requests are rejected", "Valid OIDC users can sign in"],
                "source_references": ["SRC-0001"],
                "confidence": 0.94,
                "requires_clarification": False,
                "clarification_reasons": [],
            }]})
        return json.dumps({"questions": [
            {
                "requirement_id": "REQ-001",
                "question": "Which authentication method should the product use?",
                "reason": "The source does not select an authentication standard.",
                "severity": "critical",
                "missing_field": "authentication_method",
                "recommended_answer_type": "single_select",
                "options": ["OIDC", "SAML", "Managed identity"],
                "recommended_option": "OIDC",
                "allow_custom_answer": True,
                "blocking": True,
                "source_references": ["SRC-0001"],
            },
            {
                "requirement_id": "REQ-001",
                "question": "Should audit retention exceed one year?",
                "reason": "Retention affects storage and compliance scope.",
                "severity": "medium",
                "missing_field": "audit_retention",
                "recommended_answer_type": "single_select",
                "options": ["One year", "Three years", "Seven years"],
                "recommended_option": "Three years",
                "allow_custom_answer": True,
                "blocking": False,
                "source_references": ["SRC-0001"],
            },
        ]})

    async def embed(self, texts):
        return []


@pytest.mark.asyncio
async def test_clarification_session_interrupts_and_resumes(client):
    app.dependency_overrides[get_llm_provider] = lambda: FakeProvider()
    workbook = Workbook()
    workbook.active.title = "Requirements"
    workbook.active.append(["Requirement", "Authentication is required"])
    content = io.BytesIO()
    workbook.save(content)
    content.seek(0)
    project = (await client.post("/api/v1/projects", json={"name": "Clarification Project"})).json()
    document = (await client.post(
        f"/api/v1/projects/{project['id']}/documents",
        files={"file": ("requirements.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )).json()
    assert (await client.post(f"/api/v1/documents/{document['id']}/process")).status_code == 201

    created = await client.post(f"/api/v1/projects/{project['id']}/sessions")
    assert created.status_code == 201
    session = created.json()
    assert session["status"] == "awaiting_clarification"
    assert len(session["thread_id"]) == 36

    first = (await client.get(f"/api/v1/sessions/{session['id']}/clarifications/next")).json()
    assert len(first["options"]) == 3
    answered = await client.post(f"/api/v1/clarifications/{first['id']}/answer", json={"action": "default"})
    assert answered.json() == {"session_id": session["id"], "session_status": "awaiting_clarification", "has_next": True}

    second = (await client.get(f"/api/v1/sessions/{session['id']}/clarifications/next")).json()
    assert second["required"] is False
    finished = await client.post(f"/api/v1/clarifications/{second['id']}/answer", json={"action": "skip"})
    assert finished.json()["session_status"] == "clarifications_complete"
    assert finished.json()["has_next"] is False

    generated = await client.post(f"/api/v1/sessions/{session['id']}/generate")
    assert generated.status_code == 200
    result = generated.json()
    assert result["session_status"] == "requirements_complete"
    assert result["requirements"][0]["stable_id"] == "REQ-001"
    assert result["requirements"][0]["requirement_status"] == "explicit"
    assert result["requirements"][0]["actors"] == ["User"]
    assert result["requirements"][0]["confidence"] == 0.94
    assert result["requirements"][0]["provenance"]["title"]["origin"] == "extracted"
    assert result["requirements"][0]["provenance"]["title"]["source_references"][0]["chunk_id"] == "CHUNK-001"
    assert result["requirements"][0]["source_references"] == ["requirements.xlsx | Requirements"]

    requirement_id = result["requirements"][0]["id"]
    sources_response = await client.get(f"/api/v1/backlog/requirement/{requirement_id}/sources")
    assert sources_response.status_code == 200
    sources = sources_response.json()
    assert sources["unresolved_references"] == []
    assert sources["sources"][0]["document_name"] == "requirements.xlsx"
    assert sources["sources"][0]["section"] == "Requirements"
    assert sources["sources"][0]["page"] is None
    assert "Authentication is required" in sources["sources"][0]["snippet"]

    decomposed = await client.post(f"/api/v1/sessions/{session['id']}/decompose")
    assert decomposed.status_code == 200
    decomposition = decomposed.json()
    assert decomposition["session_status"] == "decomposition_complete"
    assert decomposition["decompositions"][0]["stable_id"] == "DEC-001"
    assert decomposition["decompositions"][0]["requirement_ids"] == ["REQ-001"]
    assert decomposition["decompositions"][0]["provenance"]["title"]["origin"] == "inferred"
    assert decomposition["decompositions"][0]["provenance"]["title"]["requires_review"] is True

    backlog_response = await client.post(f"/api/v1/sessions/{session['id']}/backlog")
    assert backlog_response.status_code == 200
    backlog = backlog_response.json()
    assert backlog["session_status"] == "backlog_complete"
    assert backlog["epics"][0]["stable_id"] == "EPIC-001"
    assert backlog["stories"][0]["stable_id"] == "STORY-001"
    assert backlog["stories"][0]["epic_stable_id"] == "EPIC-001"
    assert backlog["tasks"][0]["stable_id"] == "TASK-001"
    assert backlog["tasks"][0]["story_stable_id"] == "STORY-001"
    assert backlog["tasks"][0]["requirement_ids"] == ["REQ-001"]
    assert backlog["tasks"][0]["decomposition_ids"] == ["DEC-001"]
    assert backlog["tasks"][0]["estimated_hours"] is None
    assert backlog["tasks"][0]["provenance"]["title"]["requires_review"] is True

    enriched_response = await client.post(f"/api/v1/sessions/{session['id']}/enrich")
    assert enriched_response.status_code == 200
    enriched = enriched_response.json()
    assert enriched["session_status"] == "enrichment_complete"
    assert enriched["epics"][0]["priority"] == "Critical"
    assert enriched["stories"][0]["priority"] == "High"
    assert enriched["stories"][0]["acceptance_criteria"] == [
        "Given a valid OIDC identity, when login completes, then the dashboard opens."
    ]
    assert enriched["stories"][0]["story_points"] is None
    assert enriched["tasks"][0]["estimated_hours"] is None
    assert enriched["epics"][0]["provenance"]["acceptance_criteria"]["origin"] == "inferred"

    estimated_response = await client.post(f"/api/v1/sessions/{session['id']}/estimate")
    assert estimated_response.status_code == 200
    estimated = estimated_response.json()
    assert estimated["session_status"] == "estimation_complete"
    assert estimated["stories"][0]["story_points"] == 5
    assert estimated["tasks"][0]["estimated_hours"] == 12
    assert "moderate complexity" in estimated["stories"][0]["estimation_rationale"]
    assert estimated["stories"][0]["provenance"]["story_points"]["origin"] == "calculated"
    assert estimated["tasks"][0]["provenance"]["estimated_hours"]["requires_review"] is True

    dependencies_response = await client.post(f"/api/v1/sessions/{session['id']}/dependencies")
    assert dependencies_response.status_code == 200
    dependencies = dependencies_response.json()
    assert dependencies["session_status"] == "dependencies_complete"
    assert dependencies["dependencies"] == []

    planning_workbook = Workbook()
    planning_workbook.active.title = "Teams"
    planning_workbook.active.append(["Team ID", "Team Name"])
    planning_workbook.active.append(["TEAM-1", "Platform"])
    members = planning_workbook.create_sheet("TeamMembers")
    members.append(["Team Member ID", "Team ID", "Name", "Role", "Skills", "Capacity Hours", "Allocation %"])
    members.append(["MEM-1", "TEAM-1", "Asha", "Engineer", "Python, OIDC, FastAPI", 60, 100])
    sprints = planning_workbook.create_sheet("Sprints")
    sprints.append(["Sprint ID", "Sprint Name", "Start Date", "End Date", "Capacity Points"])
    sprints.append(["SPR-1", "Sprint 1", "2026-09-21", "2026-10-02", 20])
    sprints.append(["SPR-2", "Sprint 2", "2026-10-05", "2026-10-16", 20])
    holidays = planning_workbook.create_sheet("Holidays")
    holidays.append(["Date", "Holiday Name", "Location"])
    holidays.append(["2026-09-25", "Local Holiday", "Pune"])
    planning_content = io.BytesIO()
    planning_workbook.save(planning_content)
    planning_content.seek(0)
    planning_response = await client.post(
        f"/api/v1/sessions/{session['id']}/planning-data",
        files={"file": ("Foreman_Synthetic_Dataset.xlsx", planning_content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert planning_response.status_code == 200
    assert planning_response.json()["requires_clarification"] is False
    assert planning_response.json()["members_imported"] == 1

    assignment_response = await client.post(f"/api/v1/sessions/{session['id']}/assign")
    assert assignment_response.status_code == 200
    assignment = assignment_response.json()
    assert assignment["session_status"] == "assignment_complete"
    assert len(assignment["assignments"]) == 2
    assert assignment["assignments"][1]["team_member_name"] == "Asha"
    assert assignment["assignments"][1]["recommended_hours"] == 12
    assert assignment["assignments"][1]["status"] == "proposed"

    sprint_response = await client.post(f"/api/v1/sessions/{session['id']}/plan-sprints")
    assert sprint_response.status_code == 200
    sprint_plan = sprint_response.json()
    assert sprint_plan["session_status"] == "sprint_planning_complete"
    assert sprint_plan["decisions"][0]["decision"] == "planned"
    assert sprint_plan["decisions"][0]["sprint_name"] == "Sprint 1"
    assert sprint_plan["decisions"][0]["assignee_name"] == "Asha"
    assert sprint_plan["decisions"][0]["status"] == "proposed"

    duplicate_response = await client.post(f"/api/v1/sessions/{session['id']}/duplicates")
    assert duplicate_response.status_code == 200
    assert duplicate_response.json()["session_status"] == "duplicates_complete"
    assert duplicate_response.json()["candidates"] == []

    quality_response = await client.post(f"/api/v1/sessions/{session['id']}/quality")
    assert quality_response.status_code == 200
    quality = quality_response.json()
    assert quality["session_status"] == "quality_complete"
    assert all(item["score"] >= 80 for item in quality["results"])

    health_response = await client.post(f"/api/v1/sessions/{session['id']}/board-health")
    assert health_response.status_code == 200
    health = health_response.json()
    assert health["session_status"] == "awaiting_approval"
    assert health["health"]["status"] == "review_required"
    assert health["health"]["metrics"]["overloaded_sprints"] == 0

    blocked_publish = await client.post(f"/api/v1/sessions/{session['id']}/publish")
    assert blocked_publish.status_code == 409
    assert "human approval" in blocked_publish.json()["detail"].lower()

    approval_response = await client.post(
        f"/api/v1/sessions/{session['id']}/approval",
        json={"decision": "approve", "approved_by": "Product Owner", "note": "Reviewed synthetic workflow output."},
    )
    assert approval_response.status_code == 200
    assert approval_response.json()["session_status"] == "approved"

    publish_response = await client.post(f"/api/v1/sessions/{session['id']}/publish")
    assert publish_response.status_code == 200
    published = publish_response.json()
    assert published["status"] == "completed"
    assert published["filename"] == "sprint_sarthi_backlog.xlsx"
    download = await client.get(published["download_url"])
    assert download.status_code == 200
    exported = load_workbook(io.BytesIO(download.content), read_only=True, data_only=True)
    assert exported.sheetnames == ["Epics", "User Stories", "Tasks", "Sprint Plan", "Dependencies", "Quality Report"]
    assert tuple(cell.value for cell in exported["Quality Report"][1]) == (
        "Item ID", "Item Type", "Quality Score", "Missing Acceptance Criteria", "Duplicate",
        "Ambiguous", "Dependency Issue", "Estimation Issue", "Recommendation",
    )
    story_rows = list(exported["User Stories"].iter_rows(min_row=2, values_only=True))
    sprint_rows = list(exported["Sprint Plan"].iter_rows(min_row=2, values_only=True))
    assert story_rows and all(row[0] and row[11] and row[14] for row in story_rows)
    story_sprint_rows = [row for row in sprint_rows if row[1]]
    empty_sprint_rows = [row for row in sprint_rows if not row[1]]
    assert story_sprint_rows and all(row[0] and row[1] and row[8] for row in story_sprint_rows)
    assert empty_sprint_rows == [("Sprint 2", None, None, None, None, None, None, None, "No stories planned for this sprint.")]

    usage_response = await client.get(f"/api/v1/sessions/{session['id']}/usage")
    assert usage_response.status_code == 200
    usage = usage_response.json()
    assert usage["input_tokens"] == 1100
    assert usage["output_tokens"] == 220
    assert usage["total_tokens"] == 1320
    assert usage["remaining_tokens"] == 9880
    assert [execution["node_name"] for execution in usage["executions"]] == [
        "Requirement", "Clarification", "Requirement", "Decomposition", "Backlog", "Enrichment", "Estimation",
        "Dependency", "Assignment", "Sprint", "Duplicate", "Quality", "BoardHealth"
    ]
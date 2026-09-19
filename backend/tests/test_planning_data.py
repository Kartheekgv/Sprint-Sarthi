from io import BytesIO
import json
from types import SimpleNamespace

from openpyxl import Workbook
import pytest

from app.schemas.planning_data import AssignmentBatch
from app.services.assignment import _deterministic_assignment, _rebalance_for_capacity
from app.services.planning_data import build_planning_template, parse_planning_workbook
from app.services.quality import calculate_board_health_score


def _workbook(include_skills: bool = True) -> bytes:
    workbook = Workbook()
    teams = workbook.active
    teams.title = "Teams"
    teams.append(["Team ID", "Team Name"])
    teams.append(["TEAM-1", "Platform"])
    members = workbook.create_sheet("TeamMembers")
    members.append(["Team Member ID", "Team ID", "Name", "Role", "Skills", "Capacity Hours", "Allocation %"])
    members.append(["MEM-1", "TEAM-1", "Asha", "Engineer", "Python, FastAPI" if include_skills else "", 60, 80])
    sprints = workbook.create_sheet("Sprints")
    sprints.append(["Sprint ID", "Sprint Name", "Start Date", "End Date", "Capacity Points"])
    sprints.append(["SPR-1", "Sprint 1", "2026-09-21", "2026-10-02", 20])
    holidays = workbook.create_sheet("Holidays")
    holidays.append(["Date", "Holiday Name", "Location"])
    holidays.append(["2026-09-25", "Local Holiday", "Pune"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_planning_workbook_parses_verified_fields():
    parsed = parse_planning_workbook(_workbook())
    assert not any(issue.blocking for issue in parsed.issues)
    assert parsed.members[0]["skills"] == ["Python", "FastAPI"]
    assert parsed.sprints[0]["capacity_points"] == 20
    assert parsed.holidays[0]["name"] == "Local Holiday"


def test_planning_workbook_requires_member_skills():
    parsed = parse_planning_workbook(_workbook(include_skills=False))
    assert any(issue.blocking and issue.field == "skills" for issue in parsed.issues)


def test_downloadable_planning_template_is_upload_compatible():
    content = build_planning_template()
    parsed = parse_planning_workbook(content)

    assert not any(issue.blocking for issue in parsed.issues)
    assert len(parsed.teams) == 2
    assert len(parsed.members) == 4
    assert len(parsed.sprints) == 3
    assert len(parsed.holidays) == 1
    assert len(parsed.leaves) == 1


@pytest.mark.asyncio
async def test_planning_template_download(client):
    response = await client.get("/api/v1/planning-data/template")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "Sprint-Sarthi-Planning-Data-Template.xlsx" in response.headers["content-disposition"]
    assert not any(issue.blocking for issue in parse_planning_workbook(response.content).issues)


def test_board_health_uses_normalized_rates_for_large_boards():
    score, risk_level, components = calculate_board_health_score(
        average_quality=100, quality_items=211, duplicate_candidates=13,
        dependencies=0, high_risk_dependencies=0, stories=61,
        deferred_stories=21, overloaded_sprints=0, sprints=10,
        unassigned_stories=0,
    )

    assert score == 89
    assert risk_level == "medium"
    assert components == {
        "quality": 35.0, "duplicate_readiness": 9.4,
        "dependency_readiness": 15.0, "delivery_feasibility": 19.7,
        "ownership_coverage": 10.0,
    }


def test_assignment_rebalances_overloaded_member_with_verified_capacity():
    batch = AssignmentBatch.model_validate({"assignments": [
        {"item_stable_id": "TASK-001", "team_member_id": "MEM-1", "recommended_hours": 60, "match_score": 0.9, "reason": "Initial skill-based recommendation.", "confidence": 0.9},
        {"item_stable_id": "TASK-002", "team_member_id": "MEM-1", "recommended_hours": 30, "match_score": 0.8, "reason": "Initial skill-based recommendation.", "confidence": 0.8},
    ]})
    tasks = {
        "TASK-001": SimpleNamespace(title="Build API", description="Python FastAPI service", task_type="backend", estimated_hours=60),
        "TASK-002": SimpleNamespace(title="Build UI", description="React TypeScript screen", task_type="frontend", estimated_hours=30),
    }
    members = {
        "MEM-1": SimpleNamespace(name="Asha", skills_json=json.dumps(["Python", "FastAPI"]), capacity_hours=64),
        "MEM-2": SimpleNamespace(name="Mira", skills_json=json.dumps(["React", "TypeScript"]), capacity_hours=40),
    }

    result = _rebalance_for_capacity(batch, tasks, members)

    assert result is not None
    assignments = {item.item_stable_id: item for item in result.assignments}
    assert assignments["TASK-001"].team_member_id == "MEM-1"
    assert assignments["TASK-002"].team_member_id == "MEM-2"
    assert "Capacity-adjusted" in assignments["TASK-002"].reason


def test_assignment_fallback_creates_a_reviewable_recommendation():
    task = SimpleNamespace(
        title="Build API", description="Python FastAPI service", task_type="backend",
        estimated_hours=12,
    )
    members = {
        "MEM-1": SimpleNamespace(
            name="Asha", skills_json=json.dumps(["Python", "FastAPI"]),
        ),
    }

    result = _deterministic_assignment(
        "TASK-001", {"TASK-001": task}, members,
        {"MEM-1": 40}, {},
    )

    assert result["item_stable_id"] == "TASK-001"
    assert result["team_member_id"] == "MEM-1"
    assert result["recommended_hours"] == 12
    assert result["confidence"] < 1
    assert "Human review required" in result["reason"]


@pytest.mark.asyncio
async def test_planning_upload_reports_clarification_without_inventing_data(client):
    project = (await client.post("/api/v1/projects", json={"name": "Planning Import"})).json()
    from app.core.database import get_session
    from app.main import app
    from app.models.entities import AnalysisSession

    dependency = app.dependency_overrides[get_session]
    async for db in dependency():
        session = AnalysisSession(project_id=project["id"], thread_id="11111111-1111-1111-1111-111111111111", status="dependencies_complete", current_node="Assignment")
        db.add(session)
        await db.commit()
        await db.refresh(session)
        session_id = session.id
        break
    response = await client.post(
        f"/api/v1/sessions/{session_id}/planning-data",
        files={"file": ("Foreman_Synthetic_Dataset.xlsx", _workbook(include_skills=False), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200
    result = response.json()
    assert result["requires_clarification"] is True
    assert result["members_imported"] == 0
    assert any(issue["field"] == "skills" and issue["blocking"] for issue in result["issues"])
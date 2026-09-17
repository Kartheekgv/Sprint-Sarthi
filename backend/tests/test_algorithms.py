from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.algorithms.dependencies import DependencyEdge, analyze_dependencies
from app.algorithms.quality import QualitySignals, score_quality
from app.exporters.excel import SHEETS, build_workbook


def test_dependency_analysis_orders_prerequisites_first():
    analysis = analyze_dependencies(
        {"STORY-1", "STORY-2", "STORY-3"},
        [DependencyEdge("STORY-2", "STORY-1"), DependencyEdge("STORY-3", "STORY-2")],
    )
    assert analysis.topological_order == ("STORY-1", "STORY-2", "STORY-3")
    assert not analysis.cycles


def test_dependency_analysis_rejects_self_edges_and_reports_cycles():
    with pytest.raises(ValueError, match="Self-dependency"):
        analyze_dependencies({"A"}, [DependencyEdge("A", "A")])
    analysis = analyze_dependencies({"A", "B"}, [DependencyEdge("A", "B"), DependencyEdge("B", "A")])
    assert analysis.cycles
    assert not analysis.topological_order


def test_quality_weights_total_exactly_100():
    assert score_quality(QualitySignals(True, True, True, True, True, True)) == 100
    assert score_quality(QualitySignals(False, True, True, True, True, True)) == 75


def test_workbook_matches_contract(tmp_path: Path):
    output = tmp_path / "sprint_sarthi_backlog.xlsx"
    build_workbook(output, {"Epics": [{"Epic ID": "EPIC-1", "Epic Title": "Secure intake"}]})
    workbook = load_workbook(output)
    assert workbook.sheetnames == list(SHEETS)
    for name, columns in SHEETS.items():
        sheet = workbook[name]
        assert tuple(cell.value for cell in sheet[1]) == columns
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref

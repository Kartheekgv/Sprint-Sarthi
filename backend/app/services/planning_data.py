import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.entities import Capacity, Department, Holiday, Leave, Sprint, TeamMember
from app.schemas.planning_data import PlanningIssue


def _key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _text(value: object) -> str:
    return str(value or "").strip()


def _number(value: object) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace("%", "").strip())
    except ValueError:
        return None


def _date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value:
        try:
            return datetime.fromisoformat(str(value).strip()).date()
        except ValueError:
            return None
    return None


def _split(value: object) -> list[str]:
    return [item.strip() for item in re.split(r"[,;|]", _text(value)) if item.strip()]


ALIASES = {
    "team_id": {"teamid", "departmentid", "team"},
    "team_name": {"teamname", "department", "departmentname", "team"},
    "member_id": {"teammemberid", "memberid", "employeeid", "userid", "id"},
    "member_name": {"teammember", "membername", "employeename", "name"},
    "role": {"role", "designation", "jobrole"},
    "skills": {"skills", "skillset", "primaryskills", "skill"},
    "capacity_hours": {"capacityhours", "availablehours", "capacity", "hoursavailable"},
    "allocation_percent": {"allocationpercent", "allocationpercentage", "allocation"},
    "location": {"location", "office", "region"},
    "sprint_id": {"sprintid", "iterationid", "id"},
    "sprint_name": {"sprintname", "iteration", "name"},
    "start_date": {"startdate", "sprintstart"},
    "end_date": {"enddate", "sprintend"},
    "capacity_points": {"capacitypoints", "velocity", "capacity"},
    "committed_points": {"committedpoints", "commitment", "plannedpoints"},
    "holiday_date": {"date", "holidaydate", "day"},
    "holiday_name": {"holidayname", "name", "description"},
    "leave_start": {"startdate", "leavestart", "date"},
    "leave_end": {"enddate", "leaveend"},
    "reason": {"reason", "leavetype", "description"},
}


@dataclass
class PlanningWorkbook:
    teams: list[dict[str, object]] = field(default_factory=list)
    members: list[dict[str, object]] = field(default_factory=list)
    sprints: list[dict[str, object]] = field(default_factory=list)
    holidays: list[dict[str, object]] = field(default_factory=list)
    leaves: list[dict[str, object]] = field(default_factory=list)
    issues: list[PlanningIssue] = field(default_factory=list)


def build_planning_template() -> bytes:
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    instructions.append(["Planning Data Template", "Replace the sample values before upload."])
    instructions.append(["Sheet", "Purpose"])
    instructions.append(["Teams", "Team identifiers and names."])
    instructions.append(["TeamMembers", "Skills and verified available hours used for Task assignment."])
    instructions.append(["Sprints", "Sprint dates and available story-point capacity."])
    instructions.append(["Holidays", "Location holidays used to validate capacity."])
    instructions.append(["Leaves", "Known member leave used to adjust capacity."])

    teams = workbook.create_sheet("Teams")
    teams.append(["Team ID", "Team Name"])
    teams.append(["TEAM-PLATFORM", "Platform Team"])
    teams.append(["TEAM-EXPERIENCE", "Experience Team"])

    members = workbook.create_sheet("TeamMembers")
    members.append([
        "Team Member ID", "Team ID", "Name", "Role", "Skills",
        "Capacity Hours", "Allocation %", "Location",
    ])
    members.append(["MEM-001", "TEAM-PLATFORM", "Sample Backend Engineer", "Backend Engineer", "Python, FastAPI, SQL", 80, 100, "Pune"])
    members.append(["MEM-002", "TEAM-PLATFORM", "Sample QA Engineer", "QA Engineer", "API Testing, Pytest, Automation", 64, 80, "Pune"])
    members.append(["MEM-003", "TEAM-EXPERIENCE", "Sample Frontend Engineer", "Frontend Engineer", "React, TypeScript, Accessibility", 80, 100, "Berlin"])
    members.append(["MEM-004", "TEAM-EXPERIENCE", "Sample UX Designer", "UX Designer", "UX Research, Figma, Design Systems", 48, 60, "Berlin"])

    next_monday = date.today() + timedelta(days=(7 - date.today().weekday()) % 7)
    sprints = workbook.create_sheet("Sprints")
    sprints.append(["Sprint ID", "Sprint Name", "Start Date", "End Date", "Capacity Points", "Committed Points"])
    for index in range(3):
        start = next_monday + timedelta(days=index * 14)
        sprints.append([f"SPR-{index + 1:02d}", f"Sprint {index + 1}", start, start + timedelta(days=11), 40, 0])

    holidays = workbook.create_sheet("Holidays")
    holidays.append(["Date", "Holiday Name", "Location"])
    holidays.append([next_monday + timedelta(days=4), "Sample Local Holiday", "Pune"])

    leaves = workbook.create_sheet("Leaves")
    leaves.append(["Team Member ID", "Start Date", "End Date", "Reason"])
    leaves.append(["MEM-003", next_monday + timedelta(days=7), next_monday + timedelta(days=8), "Sample planned leave"])

    header_fill = PatternFill("solid", fgColor="17324D")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _rows(sheet) -> list[tuple[int, dict[str, object]]]:
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return []
    headers = [_key(value) for value in values[0]]
    return [
        (index, {headers[column]: value for column, value in enumerate(row) if column < len(headers) and headers[column]})
        for index, row in enumerate(values[1:], start=2)
        if any(value not in (None, "") for value in row)
    ]


def _get(row: dict[str, object], field_name: str) -> object:
    return next((row[key] for key in ALIASES[field_name] if key in row), None)


def parse_planning_workbook(content: bytes) -> PlanningWorkbook:
    result = PlanningWorkbook()
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("The planning workbook is not a valid XLSX file") from error
    sheets = {_key(name): workbook[name] for name in workbook.sheetnames}
    if "teammembers" not in sheets:
        result.issues.append(PlanningIssue(sheet="TeamMembers", field="sheet", message="Required TeamMembers sheet is missing.", blocking=True))
    if "sprints" not in sheets:
        result.issues.append(PlanningIssue(sheet="Sprints", field="sheet", message="Required Sprints sheet is missing.", blocking=True))
    if "holidays" not in sheets:
        result.issues.append(PlanningIssue(sheet="Holidays", field="sheet", message="Required Holidays sheet is missing; holiday-adjusted capacity cannot be verified.", blocking=True))

    if "teams" in sheets:
        for row_number, row in _rows(sheets["teams"]):
            team_id = _text(_get(row, "team_id"))
            name = _text(_get(row, "team_name")) or team_id
            if name:
                result.teams.append({"row": row_number, "external_id": team_id or name, "name": name})
    if "teammembers" in sheets:
        for row_number, row in _rows(sheets["teammembers"]):
            values = {
                "row": row_number, "external_id": _text(_get(row, "member_id")),
                "team_id": _text(_get(row, "team_id")), "name": _text(_get(row, "member_name")),
                "role": _text(_get(row, "role")), "skills": _split(_get(row, "skills")),
                "capacity_hours": _number(_get(row, "capacity_hours")),
                "allocation_percent": _number(_get(row, "allocation_percent")),
                "location": _text(_get(row, "location")),
            }
            for field_name in ("name", "role", "skills"):
                if not values[field_name]:
                    result.issues.append(PlanningIssue(sheet="TeamMembers", row=row_number, field=field_name, message=f"{field_name.replace('_', ' ').title()} is required for assignment.", blocking=True))
            if values["capacity_hours"] is None and values["allocation_percent"] is None:
                result.issues.append(PlanningIssue(sheet="TeamMembers", row=row_number, field="availability", message="Capacity hours or allocation percent is required.", blocking=True))
            result.members.append(values)
    if "sprints" in sheets:
        for row_number, row in _rows(sheets["sprints"]):
            values = {
                "row": row_number, "external_id": _text(_get(row, "sprint_id")),
                "name": _text(_get(row, "sprint_name")), "start_date": _date(_get(row, "start_date")),
                "end_date": _date(_get(row, "end_date")), "capacity_points": _number(_get(row, "capacity_points")),
                "committed_points": _number(_get(row, "committed_points")) or 0,
            }
            for field_name in ("name", "start_date", "end_date", "capacity_points"):
                if not values[field_name]:
                    result.issues.append(PlanningIssue(sheet="Sprints", row=row_number, field=field_name, message=f"{field_name.replace('_', ' ').title()} is required for sprint planning.", blocking=True))
            if values["start_date"] and values["end_date"] and values["end_date"] < values["start_date"]:
                result.issues.append(PlanningIssue(sheet="Sprints", row=row_number, field="end_date", message="Sprint end date must not precede its start date.", blocking=True))
            result.sprints.append(values)
    if "holidays" in sheets:
        for row_number, row in _rows(sheets["holidays"]):
            day = _date(_get(row, "holiday_date"))
            name = _text(_get(row, "holiday_name"))
            if not day or not name:
                result.issues.append(PlanningIssue(sheet="Holidays", row=row_number, field="date/name", message="Holiday date and name are required.", blocking=True))
            result.holidays.append({"row": row_number, "day": day, "name": name, "location": _text(_get(row, "location"))})
    leave_sheet = sheets.get("leaves") or sheets.get("leavecalendar")
    if leave_sheet:
        for row_number, row in _rows(leave_sheet):
            result.leaves.append({
                "row": row_number, "member_id": _text(_get(row, "member_id")),
                "start_date": _date(_get(row, "leave_start")),
                "end_date": _date(_get(row, "leave_end")) or _date(_get(row, "leave_start")),
                "reason": _text(_get(row, "reason")),
            })
    else:
        result.issues.append(PlanningIssue(sheet="Leaves", field="sheet", message="No leave sheet was supplied; member capacity must already be leave-adjusted.", blocking=False))
    return result


async def persist_planning_workbook(
    db: AsyncSession, project_id: str, filename: str, parsed: PlanningWorkbook
) -> None:
    if any(issue.blocking for issue in parsed.issues):
        return
    await db.execute(delete(Capacity).where(Capacity.sprint_id.in_(select(Sprint.id).where(Sprint.project_id == project_id))))
    await db.execute(delete(Leave).where(Leave.project_id == project_id))
    await db.execute(delete(Holiday).where(Holiday.project_id == project_id))
    await db.execute(delete(Sprint).where(Sprint.project_id == project_id))
    await db.execute(delete(TeamMember).where(TeamMember.project_id == project_id))
    await db.execute(delete(Department).where(Department.project_id == project_id))
    provenance = lambda sheet, row: json.dumps({"origin": "human_provided", "file_name": filename, "sheet": sheet, "row": row})
    teams = parsed.teams or [{"row": 0, "external_id": "UNASSIGNED", "name": "Unassigned Team"}]
    department_by_external: dict[str, Department] = {}
    for item in teams:
        record = Department(project_id=project_id, external_id=str(item["external_id"]), name=str(item["name"]))
        db.add(record)
        department_by_external[str(item["external_id"])] = record
    await db.flush()
    members_by_external: dict[str, TeamMember] = {}
    for item in parsed.members:
        department = department_by_external.get(str(item["team_id"])) or next(iter(department_by_external.values()))
        record = TeamMember(
            project_id=project_id, department_id=department.id,
            external_id=str(item["external_id"] or item["name"]), name=str(item["name"]),
            role=str(item["role"]), skills_json=json.dumps(item["skills"]),
            capacity_hours=item["capacity_hours"], allocation_percent=item["allocation_percent"] or 100,
            location=str(item["location"]), provenance_json=provenance("TeamMembers", item["row"]),
        )
        db.add(record)
        members_by_external[record.external_id or record.name] = record
    await db.flush()
    for item in parsed.sprints:
        db.add(Sprint(
            project_id=project_id, external_id=str(item["external_id"] or item["name"]),
            name=str(item["name"]), start_date=item["start_date"], end_date=item["end_date"],
            capacity_points=item["capacity_points"], committed_points=item["committed_points"],
            provenance_json=provenance("Sprints", item["row"]),
        ))
    for item in parsed.holidays:
        db.add(Holiday(
            project_id=project_id, day=item["day"], name=str(item["name"]),
            location=str(item["location"]), provenance_json=provenance("Holidays", item["row"]),
        ))
    for item in parsed.leaves:
        member = members_by_external.get(str(item["member_id"]))
        if member and item["start_date"] and item["end_date"]:
            db.add(Leave(
                project_id=project_id, team_member_id=member.id,
                start_date=item["start_date"], end_date=item["end_date"], reason=str(item["reason"]),
                provenance_json=provenance("Leaves", item["row"]),
            ))
    await db.flush()
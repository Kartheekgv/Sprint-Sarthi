import csv
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document as WordDocument
from openpyxl import load_workbook


MAX_CHUNK_CHARS = 4000


class DocumentExtractionError(RuntimeError):
    pass


@dataclass(frozen=True)
class ExtractedSection:
    section: str
    heading: str
    page: int | None
    content: str


def _chunks(text: str) -> list[str]:
    normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    return [normalized[index:index + MAX_CHUNK_CHARS] for index in range(0, len(normalized), MAX_CHUNK_CHARS)]


def _extract_pdf(path: Path) -> list[ExtractedSection]:
    sections: list[ExtractedSection] = []
    with pymupdf.open(path) as document:
        for page_index, page in enumerate(document):
            for content in _chunks(page.get_text("text")):
                sections.append(ExtractedSection("Document", f"Page {page_index + 1}", page_index + 1, content))
    return sections


def _extract_docx(path: Path) -> list[ExtractedSection]:
    document = WordDocument(path)
    blocks: list[tuple[str, str]] = []
    heading = "Document"
    paragraphs: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name.lower().startswith("heading"):
            if paragraphs:
                blocks.append((heading, "\n".join(paragraphs)))
                paragraphs = []
            heading = text
        else:
            paragraphs.append(text)
    if paragraphs:
        blocks.append((heading, "\n".join(paragraphs)))
    for table_index, table in enumerate(document.tables, 1):
        rows = ["\t".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        blocks.append((f"Table {table_index}", "\n".join(rows)))
    return [ExtractedSection("Document", title, None, chunk) for title, text in blocks for chunk in _chunks(text)]


def _extract_xlsx(path: Path) -> list[ExtractedSection]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[ExtractedSection] = []
    try:
        for sheet in workbook.worksheets:
            rows = ["\t".join("" if value is None else str(value) for value in row) for row in sheet.iter_rows(values_only=True)]
            for content in _chunks("\n".join(rows)):
                sections.append(ExtractedSection(sheet.title, sheet.title, None, content))
    finally:
        workbook.close()
    return sections


def _extract_text(path: Path) -> list[ExtractedSection]:
    text = path.read_text(encoding="utf-8-sig")
    if path.suffix.lower() != ".md":
        return [ExtractedSection("Document", "Document", None, chunk) for chunk in _chunks(text)]
    sections: list[ExtractedSection] = []
    heading = "Document"
    content: list[str] = []
    for line in text.splitlines():
        if line.lstrip().startswith("#"):
            if content:
                sections.extend(ExtractedSection("Document", heading, None, chunk) for chunk in _chunks("\n".join(content)))
                content = []
            heading = line.lstrip("#").strip() or "Document"
        else:
            content.append(line)
    if content:
        sections.extend(ExtractedSection("Document", heading, None, chunk) for chunk in _chunks("\n".join(content)))
    return sections


def _extract_csv(path: Path) -> list[ExtractedSection]:
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        rows = ["\t".join(row) for row in csv.reader(source)]
    return [ExtractedSection("CSV", "CSV", None, chunk) for chunk in _chunks("\n".join(rows))]


def extract_document(path: Path) -> list[ExtractedSection]:
    try:
        extractor = {
            ".pdf": _extract_pdf,
            ".docx": _extract_docx,
            ".xlsx": _extract_xlsx,
            ".txt": _extract_text,
            ".md": _extract_text,
            ".csv": _extract_csv,
        }.get(path.suffix.lower())
        if extractor is None:
            raise DocumentExtractionError("Unsupported document type")
        sections = extractor(path)
    except DocumentExtractionError:
        raise
    except Exception as error:
        raise DocumentExtractionError("The document could not be parsed") from error
    if not sections:
        raise DocumentExtractionError("The document contains no extractable text")
    return sections
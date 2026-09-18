from collections.abc import Mapping, Sequence
from datetime import date, datetime
from itertools import islice
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SHEETS: dict[str, tuple[str, ...]] = {
    "Epics": ("Epic ID", "Architecture Layer", "Epic Title", "Description", "Business Value", "Priority", "Acceptance Criteria", "Source Reference", "Quality Score"),
    "User Stories": ("Story ID", "Feature ID", "Feature Title", "Epic ID", "Story Title", "User Story", "Description", "Acceptance Criteria", "Definition of Done", "Priority", "Story Points", "Dependencies", "Suggested Assignee", "Department", "Sprint", "Status", "Quality Score", "Source Reference"),
    "Tasks": ("Task ID", "Story ID", "Feature ID", "Epic ID", "Task Title", "Description", "Task Type", "Work Category", "Acceptance Criteria", "Definition of Done", "Priority", "Estimated Hours", "Dependencies", "Suggested Assignee", "Department", "Sprint", "Status", "Source Reference"),
    "Sprint Plan": ("Sprint", "Story ID", "Story Title", "Story Points", "Suggested Assignee", "Available Capacity", "Dependencies", "Priority", "Reason for Selection"),
    "Dependencies": ("Source ID", "Source Title", "Target ID", "Target Title", "Dependency Type", "Risk", "Explanation"),
    "Quality Report": ("Item ID", "Item Type", "Quality Score", "Missing Acceptance Criteria", "Duplicate", "Ambiguous", "Dependency Issue", "Estimation Issue", "Recommendation"),
}

REQUIRED_SHEETS = tuple(SHEETS)


def build_workbook(destination: Path, rows: Mapping[str, Sequence[Mapping[str, object]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="D95500")
    header_font = Font(color="FFFFFF", bold=True)

    for sheet_name, columns in SHEETS.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(columns)
        for record in rows.get(sheet_name, []):
            sheet.append([record.get(column, "") for column in columns])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 26
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column_index, column in enumerate(columns, 1):
            values = [str(sheet.cell(row=row, column=column_index).value or "") for row in range(1, sheet.max_row + 1)]
            sheet.column_dimensions[get_column_letter(column_index)].width = min(max(max(map(len, values)) + 2, 12), 50)
            for cell in list(sheet.columns)[column_index - 1][1:]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)


def validate_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    if tuple(workbook.sheetnames) != REQUIRED_SHEETS:
        raise ValueError("Published workbook does not contain the required 6 sheets in order")
    for sheet_name, columns in SHEETS.items():
        sheet = workbook[sheet_name]
        headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
        if headers != columns:
            raise ValueError(f"Published workbook has invalid headers in {sheet_name}")


def preview_workbook(path: Path, sheet_name: str | None, offset: int, limit: int) -> dict[str, object]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_names = list(workbook.sheetnames)
    selected_name = sheet_name or sheet_names[0]
    if selected_name not in sheet_names:
        raise ValueError(f"Unknown workbook sheet: {selected_name}")
    sheet = workbook[selected_name]
    values = sheet.iter_rows(values_only=True)
    columns = [str(value or "") for value in next(values, ())]

    def serialize(value: object) -> object:
        return value.isoformat() if isinstance(value, (date, datetime)) else value

    rows = [[serialize(value) for value in row] for row in islice(values, offset, offset + limit)]
    return {
        "sheet_names": sheet_names,
        "sheet_name": selected_name,
        "columns": columns,
        "rows": rows,
        "offset": offset,
        "limit": limit,
        "total_rows": max(sheet.max_row - 1, 0),
    }
